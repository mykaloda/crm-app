#!/usr/bin/env python3
"""All pulled CRM tables in one Excel workbook: data/export/lerega-crm.xlsx.

One sheet per table plus a data dictionary from the CRM's describe_data.
Times are converted to Austin time. Needs openpyxl (pip install openpyxl).
"""

import json
import os
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
TZ = ZoneInfo("America/Chicago")
TABLES = ("orders", "stage_history", "payments", "estimates", "appointments", "work_time",
          "notes", "messages", "site_visits", "stages")


def data_dir():
    return Path(os.environ.get("CRM_DATA_DIR") or HERE / "data")


def cell_value(value, column_type):
    if value is None:
        return None
    if column_type.startswith("timestamp") and isinstance(value, str):
        t = datetime.fromisoformat(value)
        if t.utcoffset() == timedelta(0) and (t.hour, t.minute, t.second, t.microsecond) == (0, 0, 0, 0):
            return t.replace(tzinfo=None)  # a Monday.com date stored as UTC midnight: keep the date
        return t.astimezone(TZ).replace(tzinfo=None)
    if column_type == "date" and isinstance(value, str):
        return date.fromisoformat(value[:10])
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def main():
    out = data_dir() / "export" / "lerega-crm.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    described = {v["view"]: v for v in json.loads((data_dir() / "describe_data.json").read_text())}
    conn = sqlite3.connect(data_dir() / "crm.sqlite")
    present = {e for (e,) in conn.execute("SELECT DISTINCT entity FROM _records")}
    wb = Workbook(write_only=True)
    bold = Font(bold=True)  # defined before the first sheet uses it

    about = wb.create_sheet("Описание")
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 24
    about.column_dimensions["C"].width = 26
    about.column_dimensions["D"].width = 90
    head = []
    for v in ("Таблица", "Колонка", "Тип", "Что это"):
        cell = WriteOnlyCell(about, value=v)
        cell.font = bold
        head.append(cell)
    about.append(head)
    for name in TABLES:
        view = described.get(name)
        if not view:
            continue
        about.append([name, None, None, view.get("note")])
        for c in view["columns"]:
            about.append([name, c["name"], c["type"], c.get("note")])

    counts = {}
    for name in TABLES:
        if name not in present or name not in described:
            continue
        columns = described[name]["columns"]
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        for i, c in enumerate(columns, 1):
            wide = c["type"] == "text" and c["name"] in ("body", "client_request", "note", "estimator_note", "quote_text", "address")
            ws.column_dimensions[get_column_letter(i)].width = 60 if wide else max(12, min(28, len(c["name"]) + 4))
        header = []
        for c in columns:
            cell = WriteOnlyCell(ws, value=c["name"] + (" (Остин)" if c["type"].startswith("timestamp") else ""))
            cell.font = bold
            header.append(cell)
        ws.append(header)
        rows = conn.execute(
            "SELECT data FROM _records WHERE entity = ? AND deleted_at IS NULL ORDER BY rowid", (name,))  # pull order
        n = 0
        for (data,) in rows:
            record = json.loads(data)
            ws.append([cell_value(record.get(c["name"]), c["type"]) for c in columns])
            n += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n + 1}"
        counts[name] = n
    conn.close()
    wb.save(out)
    print(f"{out}: {out.stat().st_size // 1024} KiB", counts)


if __name__ == "__main__":
    main()
